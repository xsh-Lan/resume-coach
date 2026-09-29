#!/usr/bin/env python3
"""Render Markdown reports into readable HTML and DOCX files.

Markdown remains the source of truth. HTML is dependency-free; DOCX requires
python-docx. PDF is optional and can be produced by opening DOCX/HTML in Word,
WPS, or a browser and using Print / Save as PDF.
"""
from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path

CSS = """
* { box-sizing: border-box; }
:root { color-scheme: light; --blue:#2563eb; --ink:#172033; --muted:#64748b; --line:#e5e7eb; --bg:#eef2f8; }
body { margin:0; background:var(--bg); color:var(--ink); font-family:"Microsoft YaHei","PingFang SC","Segoe UI",Arial,sans-serif; }
.topbar { position:sticky; top:0; z-index:20; background:#ffffff; border-bottom:1px solid var(--line); padding:14px 24px; display:flex; align-items:center; gap:16px; box-shadow:0 2px 8px rgba(15,23,42,.05); }
.topbar .brand { font-weight:800; color:#0f172a; font-size:16px; }
.topbar .subtitle { color:var(--muted); font-size:13px; }
.layout { max-width:1200px; margin:24px auto; display:grid; grid-template-columns:264px minmax(0,1fr); gap:22px; padding:0 20px; }
.toc { position:sticky; top:84px; align-self:start; background:#fff; border:1px solid var(--line); border-radius:14px; padding:16px; max-height:calc(100vh - 104px); overflow:auto; }
.toc .toc-title { font-size:12px; letter-spacing:.08em; color:var(--muted); text-transform:uppercase; margin:0 0 8px; }
.toc a { display:block; padding:6px 8px; margin:2px 0; color:#475569; text-decoration:none; border-radius:7px; font-size:13px; line-height:1.45; }
.toc a:hover { background:#eff6ff; color:var(--blue); }
.toc a.l2 { font-weight:600; }
.toc a.l3 { padding-left:18px; }
.content { min-width:0; background:#fff; border:1px solid var(--line); border-radius:18px; padding:44px 52px; box-shadow:0 12px 34px rgba(15,23,42,.08); }
h1,h2,h3,h4 { color:#0f172a; line-height:1.3; scroll-margin-top:90px; }
h1 { font-size:28px; margin:0 0 22px; padding-bottom:16px; border-bottom:2px solid var(--blue); }
h2 { font-size:21px; margin:2.4em 0 .9em; padding-bottom:10px; border-bottom:1px solid var(--line); }
h3 { font-size:16px; margin:1.6em 0 .6em; color:#1d4ed8; }
h4 { font-size:15px; margin:1.2em 0 .4em; color:#334155; }
p,li,td,th { font-size:14.5px; line-height:1.75; }
ul,ol { padding-left:24px; }
blockquote { margin:16px 0; padding:12px 18px; background:#f8fafc; border-left:4px solid var(--blue); color:#334155; border-radius:8px; }
table { width:100%; border-collapse:collapse; margin:18px 0; font-size:13.5px; }
th { background:#f1f5f9; color:#0f172a; text-align:left; }
th,td { border:1px solid #dbe2ea; padding:9px 10px; vertical-align:top; }
tbody tr:nth-child(even) { background:#fafbfd; }
code { background:#f1f5f9; padding:2px 5px; border-radius:5px; font-family:Consolas,monospace; font-size:12.5px; color:#334155; }
pre { background:#0f172a; color:#e2e8f0; padding:16px; border-radius:10px; overflow:auto; }
pre code { background:transparent; color:inherit; padding:0; }
hr { border:0; border-top:1px solid var(--line); margin:30px 0; }
.kv { display:grid; grid-template-columns:118px 1fr; gap:6px 12px; margin:7px 0; padding:8px 11px; border:1px solid #eef2f7; border-left:3px solid var(--blue); background:#fbfdff; border-radius:7px; }
.kv .k { font-weight:700; color:#1e3a8a; white-space:nowrap; font-size:13.5px; }
.kv .v p { margin:0; }
.sev { display:inline-block; padding:2px 9px; border-radius:999px; font-size:12px; font-weight:700; }
.sev.high { background:#fee2e2; color:#b91c1c; }
.sev.medium { background:#ffedd5; color:#c2410c; }
.sev.low { background:#e5e7eb; color:#475569; }
.reading-guide { display:flex; gap:12px; align-items:flex-start; margin:0 0 22px; padding:14px 16px; background:#eef6ff; border:1px solid #c7dbff; border-radius:12px; color:#1e3a8a; font-size:14px; line-height:1.6; }
.reading-guide strong { white-space:nowrap; }
.next-steps { margin-top:36px; padding:18px 20px; background:#f8fafc; border:1px solid #e2e8f0; border-radius:12px; }
.next-steps h2 { margin-top:0; border-bottom:0; padding-bottom:0; }
.next-steps p { margin:8px 0 0; }
@media (max-width: 900px) { .layout{grid-template-columns:1fr} .toc{position:static;max-height:none} .content{padding:26px 20px} table{display:block;overflow-x:auto;white-space:nowrap} .reading-guide{flex-direction:column;gap:4px} }
@media print { body{background:#fff} .topbar,.toc{display:none} .layout{display:block;max-width:none;margin:0;padding:0} .content{border:0;box-shadow:none;border-radius:0;padding:0} h2{break-after:avoid} table{break-inside:auto} tr{break-inside:avoid} }
"""


def inline_html(s: str) -> str:
    s=html.escape(s, quote=False)
    s=re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', s)
    s=re.sub(r'`(.+?)`', r'<code>\1</code>', s)
    return s

def inline_text(s: str) -> str:
    s=re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    s=re.sub(r'`(.+?)`', r'\1', s)
    return s.replace('\t',' ')

def is_table_start(lines, i):
    return i+1 < len(lines) and lines[i].strip().startswith('|') and re.match(r'^\s*\|?[\s:|-]+\|?\s*$', lines[i+1])

def parse_table(lines, i):
    rows=[]
    while i < len(lines) and lines[i].strip().startswith('|'):
        raw=lines[i].strip().strip('|')
        cells=[inline_html(c.strip()) for c in raw.split('|')]
        if not re.match(r'^[\s:|-]+$', raw):
            rows.append(cells)
        i += 1
    return rows, i

def markdown_to_html(md: str):
    lines=md.splitlines()
    out=[]; i=0; list_type=None; in_code=False; code=[]
    headings=[]; heading_no=0
    def close_list():
        nonlocal list_type
        if list_type:
            out.append(f'</{list_type}>')
            list_type=None
    while i < len(lines):
        line=lines[i]
        if line.strip().startswith('```'):
            if in_code:
                out.append('<pre><code>'+html.escape('\n'.join(code))+'</code></pre>')
                code=[]; in_code=False
            else:
                close_list(); in_code=True
            i+=1; continue
        if in_code:
            code.append(line); i+=1; continue
        if not line.strip():
            close_list(); i+=1; continue
        if is_table_start(lines,i):
            close_list(); rows,i=parse_table(lines,i)
            if rows:
                out.append('<table><thead><tr>'+''.join(f'<th>{c}</th>' for c in rows[0])+'</tr></thead><tbody>')
                for row in rows[1:]: out.append('<tr>'+''.join(f'<td>{c}</td>' for c in row)+'</tr>')
                out.append('</tbody></table>')
            continue
        m=re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            close_list(); level=len(m.group(1)); heading_no+=1; hid=f"h{heading_no}"; headings.append((level,inline_text(m.group(2)),hid)); out.append(f'<h{level} id="{hid}">{inline_html(m.group(2))}</h{level}>'); i+=1; continue
        m=re.match(r'^\s*[-*]\s*(原句|诊断|问题等级|建议改法|可直接替换|待补信息)[：:](.*)$', line)
        if m:
            close_list()
            val=inline_html(m.group(2).strip())
            if m.group(1)=='问题等级':
                cls='high' if '高' in m.group(2) else ('medium' if '中' in m.group(2) else 'low')
                val=f'<span class="sev {cls}">{val}</span>'
            out.append('<div class="kv"><span class="k">'+inline_html(m.group(1))+'</span><div class="v"><p>'+val+'</p></div></div>')
            i+=1; continue
        m=re.match(r'^\s*[-*]\s+(.*)$', line)
        if m:
            if list_type!='ul': close_list(); out.append('<ul>'); list_type='ul'
            out.append('<li>'+inline_html(m.group(1))+'</li>'); i+=1; continue
        m=re.match(r'^\s*\d+[.)]\s+(.*)$', line)
        if m:
            if list_type!='ol': close_list(); out.append('<ol>'); list_type='ol'
            out.append('<li>'+inline_html(m.group(1))+'</li>'); i+=1; continue
        if line.startswith('>'):
            close_list(); out.append('<blockquote>'+inline_html(line.lstrip('> ').strip())+'</blockquote>'); i+=1; continue
        if line.strip() in ('---','***'):
            close_list(); out.append('<hr>'); i+=1; continue
        close_list(); out.append('<p>'+inline_html(line.strip())+'</p>'); i+=1
    close_list()
    return '\n'.join(out), headings

def markdown_to_docx(md: str, path: Path) -> tuple[bool,str]:
    try:
        from docx import Document
        from docx.shared import Pt, Inches
    except Exception as exc:
        return False, f'python-docx not installed: {exc}'
    doc=Document()
    for section in doc.sections:
        section.top_margin=Inches(0.65); section.bottom_margin=Inches(0.65)
        section.left_margin=Inches(0.72); section.right_margin=Inches(0.72)
    style=doc.styles['Normal']; style.font.name='Microsoft YaHei'; style.font.size=Pt(10.5)
    lines=md.splitlines(); i=0; in_code=False; code=[]
    while i < len(lines):
        line=lines[i]
        if line.strip().startswith('```'):
            if in_code:
                p=doc.add_paragraph('\n'.join(code)); p.style=doc.styles['No Spacing']
                for run in p.runs: run.font.name='Consolas'; run.font.size=Pt(9)
                code=[]; in_code=False
            else: in_code=True
            i+=1; continue
        if in_code: code.append(line); i+=1; continue
        if not line.strip(): i+=1; continue
        if is_table_start(lines,i):
            rows,i=parse_table(lines,i)
            if rows:
                table=doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                try: table.style='Light Grid Accent 1'
                except Exception: table.style='Table Grid'
                for ri,row in enumerate(rows):
                    for ci,cell in enumerate(row):
                        table.cell(ri,ci).text=re.sub(r'<[^>]+>','',cell)
            continue
        m=re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            doc.add_heading(inline_text(m.group(2)), level=min(len(m.group(1)),4)); i+=1; continue
        m=re.match(r'^\s*[-*]\s+(.*)$', line)
        if m:
            p=doc.add_paragraph(style='List Bullet'); p.add_run(inline_text(m.group(1))); i+=1; continue
        m=re.match(r'^\s*\d+[.)]\s+(.*)$', line)
        if m:
            p=doc.add_paragraph(style='List Number'); p.add_run(inline_text(m.group(1))); i+=1; continue
        if line.startswith('>'):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25); p.add_run(inline_text(line.lstrip('> ').strip())).italic=True; i+=1; continue
        if line.strip() in ('---','***'):
            doc.add_paragraph('────────────────────'); i+=1; continue
        doc.add_paragraph(inline_text(line.strip())); i+=1
    doc.save(path)
    return True,''

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('input', nargs='+')
    ap.add_argument('--output-dir', default=None)
    ap.add_argument('--formats', nargs='+', default=['html'], choices=['html','docx'])
    args=ap.parse_args()
    for raw in args.input:
        src=Path(raw)
        out=Path(args.output_dir) if args.output_dir else src.parent
        out.mkdir(parents=True,exist_ok=True)
        md=src.read_text(encoding='utf-8')
        if 'html' in args.formats:
            title=next((inline_text(x.lstrip('# ').strip()) for x in md.splitlines() if x.startswith('# ')),src.stem)
            body,headings=markdown_to_html(md)
            toc=[]
            for level,text,hid in headings:
                if level in (2,3):
                    toc.append(f'<a class="l{level}" href="#{hid}">{html.escape(text)}</a>')
            nav=f'<nav class="toc"><div class="toc-title">目录</div>{"".join(toc)}</nav>' if toc else ''
            top=f'<div class="topbar"><span class="brand">简历雷达</span><span class="subtitle">{html.escape(title)}</span></div>'
            guide='<div class="reading-guide"><strong>阅读指引</strong><span>先看「1. 总体结论」和「8. 10 个最优先修改点」，再按左侧目录查看逐句体检。报告只基于你提供的信息，未做外部验证；所有「待补」都是缺失信息，没有编造事实。</span></div>'
            next_steps='<div class="next-steps"><h2>下一步可以做什么</h2><p>如需进一步使用简历雷达，你可以：提供具体 JD 做正式匹配评分；补充项目原始数据、看板或代码链接用于核对证据；指定目标岗位方向生成定制改写；或继续输出 STAR 面试包。</p></div>'
            page=f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)}</title><style>{CSS}</style></head><body>{top}<div class="layout">{nav}<article class="content">{guide}{body}{next_steps}</article></div></body></html>'

            p=out/(src.stem+'.html'); p.write_text(page,encoding='utf-8'); print('HTML',p)
        if 'docx' in args.formats:
            p=out/(src.stem+'.docx'); ok,msg=markdown_to_docx(md,p)
            print('DOCX',p if ok else 'SKIPPED '+msg)
    return 0

if __name__=='__main__': raise SystemExit(main())